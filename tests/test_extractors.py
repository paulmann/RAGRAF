"""Проверка извлечения документов (G5).

Свойства:
  1. текст извлекается и содержит маркеры страницы/листа/слайда — по ним
     `path:line` остаётся осмысленным;
  2. таблица DOCX и значения XLSX попадают в текст, а не теряются;
  3. вложения EML перечисляются ИМЕНАМИ, содержимое не читается;
  4. скан без текстового слоя даёт пустой текст и причину, а не выдумку;
  5. старый двоичный формат и неизвестное расширение дают `unavailable`
     с причиной, а не исключение;
  6. нет библиотеки — `unavailable`, а не падение (проверяется подменой
     импорта).
"""

from __future__ import annotations

import builtins
import pathlib
import sys
import tempfile

REPO = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "src"))

from ragraf import documents  # noqa: E402

PASS = 0
FAIL = 0


def check(cond: bool, label: str, detail: str = "") -> None:
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  [OK ] {label}")
    else:
        FAIL += 1
        print(f"  [FAIL] {label} {detail}")


print("=" * 78)
print("1. НАЛИЧИЕ БИБЛИОТЕК")
print("=" * 78)
avail = documents.available()
check(all(avail.get(k) for k in ("pdf", "docx", "xlsx", "pptx", "eml")),
      "все извлекатели доступны", str(avail))

with tempfile.TemporaryDirectory(prefix="ragraf-docs-") as tmp:
    base = pathlib.Path(tmp)

    print()
    print("=" * 78)
    print("2. DOCX: АБЗАЦЫ И ТАБЛИЦА")
    print("=" * 78)
    import docx

    doc = docx.Document()
    doc.add_paragraph("HELLO-DOCX")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "A"
    table.rows[0].cells[1].text = "B"
    docx_path = base / "sample.docx"
    doc.save(str(docx_path))
    got = documents.extract(docx_path)
    check(got.status == "ok", "DOCX разобран", got.reason)
    check("HELLO-DOCX" in got.text, "абзац извлечён")
    check("A | B" in got.text, "строка таблицы извлечена", got.text[:120])

    print()
    print("=" * 78)
    print("3. XLSX: ЛИСТ И ЗНАЧЕНИЯ")
    print("=" * 78)
    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Данные"
    sheet.append(["HELLO-XLSX", 42])
    xlsx_path = base / "sample.xlsx"
    book.save(str(xlsx_path))
    got = documents.extract(xlsx_path)
    check(got.status == "ok", "XLSX разобран", got.reason)
    check("--- sheet: Данные ---" in got.text, "маркер листа есть")
    check("HELLO-XLSX" in got.text and "42" in got.text, "значения извлечены", got.text[:120])

    print()
    print("=" * 78)
    print("4. PPTX: СЛАЙД И ЗАМЕТКИ")
    print("=" * 78)
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.text = "HELLO-PPTX"
    pptx_path = base / "sample.pptx"
    deck.save(str(pptx_path))
    got = documents.extract(pptx_path)
    check(got.status == "ok", "PPTX разобран", got.reason)
    check("--- slide 1 ---" in got.text, "маркер слайда есть")
    check("HELLO-PPTX" in got.text, "текст слайда извлечён", got.text[:120])

    print()
    print("=" * 78)
    print("5. EML: ЗАГОЛОВКИ, ТЕЛО, ИМЕНА ВЛОЖЕНИЙ")
    print("=" * 78)
    from email.message import EmailMessage

    message = EmailMessage()
    message["Subject"] = "Тема письма"
    message["From"] = "a@example.com"
    message["To"] = "b@example.com"
    message.set_content("HELLO-EML body")
    message.add_attachment(b"secret-bytes", maintype="application", subtype="octet-stream",
                           filename="report.bin")
    eml_path = base / "sample.eml"
    eml_path.write_bytes(message.as_bytes())
    got = documents.extract(eml_path)
    check(got.status == "ok", "EML разобран", got.reason)
    check("Subject: Тема письма" in got.text, "тема извлечена")
    check("HELLO-EML body" in got.text, "тело извлечено")
    check("report.bin" in got.text, "имя вложения извлечено")
    check("secret-bytes" not in got.text, "содержимое вложения НЕ попало в текст")

    print()
    print("=" * 78)
    print("6. PDF: СТРАНИЦА БЕЗ ТЕКСТОВОГО СЛОЯ")
    print("=" * 78)
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    def text_pdf(path: pathlib.Path, label: str) -> None:
        writer = PdfWriter()
        page = writer.add_blank_page(width=200, height=200)
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 20 100 Td ({label}) Tj ET".encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(stream)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                 NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject(
                {NameObject("/F1"): writer._add_object(font)}
            )
        })
        with open(path, "wb") as handle:
            writer.write(handle)

    text_path = base / "text.pdf"
    text_pdf(text_path, "HELLO-PDF-TEXT")
    got = documents.extract(text_path)
    check(got.status == "ok", "PDF с текстом открыт", got.reason)
    check("HELLO-PDF-TEXT" in got.text, "текст PDF извлечён", got.text[:80])

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    pdf_path = base / "blank.pdf"
    with open(pdf_path, "wb") as handle:
        writer.write(handle)
    got = documents.extract(pdf_path)
    check(got.status == "ok", "пустой PDF открыт", got.reason)
    check("--- page 1 ---" in got.text, "маркер страницы есть")
    check("текстового слоя нет" in got.reason, "отсутствие текста названо", got.reason)

    print()
    print("=" * 78)
    print("6a. ODT И RTF")
    print("=" * 78)
    import zipfile

    odt_path = base / "sample.odt"
    content = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"'
        ' xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0">'
        '<office:body><office:text><text:p>HELLO-ODT</text:p>'
        '<text:p>второй абзац</text:p></office:text></office:body>'
        '</office:document-content>'
    )
    with zipfile.ZipFile(odt_path, "w") as archive:
        archive.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        archive.writestr("content.xml", content)
    got = documents.extract(odt_path)
    check(got.status == "ok", "ODT разобран", got.reason)
    check("HELLO-ODT" in got.text and "второй абзац" in got.text,
          "абзацы ODT извлечены", got.text[:120])

    rtf_path = base / "sample.rtf"
    rtf_path.write_bytes(b"{\\rtf1\\ansi HELLO-RTF\\par second line}")
    got = documents.extract(rtf_path)
    check(got.status == "ok", "RTF разобран", got.reason)
    check("HELLO-RTF" in got.text and "second line" in got.text,
          "текст RTF извлечён", got.text[:120])
    check("приблизительно" in got.reason, "грубость RTF названа", got.reason)

    print()
    print("=" * 78)
    print("6b. MSG: БЕЗ БИБЛИОТЕКИ — ЧЕСТНЫЙ ОТКАЗ")
    print("=" * 78)
    msg_path = base / "sample.msg"
    msg_path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1dummy")
    got = documents.extract(msg_path)
    check(got.status in ("unavailable", "error"),
          "MSG не выдаёт себя за разобранный", got.status)
    check(bool(got.reason), "причина MSG названа", got.reason)

    print()
    print("=" * 78)
    print("7. ЧЕСТНАЯ ДЕГРАДАЦИЯ")
    print("=" * 78)
    legacy = base / "old.doc"
    legacy.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1")
    got = documents.extract(legacy)
    check(got.status == "unavailable", "старый .doc — unavailable", got.status)
    # Контракт изменился осознанно. Раньше .doc объявлялся «нужен конвертер»,
    # потому что разборщика не было вовсе. Теперь его читает собственный
    # разборщик OLE (officebinary.py), и причина отказа называет контейнер.
    # Проверка держит новое поведение: битый контейнер не выдаёт себя за
    # документ без текста — это разные состояния.
    check("OLE" in got.reason, "причина называет контейнер", got.reason)

    unknown = base / "data.bin"
    unknown.write_bytes(b"\x00\x01\x02")
    got = documents.extract(unknown)
    check(got.status == "unavailable", "неизвестный формат — unavailable")
    check("не поддержан" in got.reason, "причина названа", got.reason)

    # Подмена импорта: библиотеки «нет» — модуль обязан вернуть unavailable.
    real_import = builtins.__import__

    def no_pypdf(name, *args, **kwargs):
        if name == "pypdf":
            raise ImportError("simulated absence")
        return real_import(name, *args, **kwargs)

    builtins.__import__ = no_pypdf
    try:
        got = documents.extract(pdf_path)
    finally:
        builtins.__import__ = real_import
    check(got.status == "unavailable", "без pypdf — unavailable, не исключение", got.status)
    check("pypdf" in got.reason, "подсказка про установку есть", got.reason)

print()
print("=" * 78)
print(f"ИТОГ: пройдено {PASS}, провалено {FAIL}")
print("=" * 78)
raise SystemExit(0 if FAIL == 0 else 1)
